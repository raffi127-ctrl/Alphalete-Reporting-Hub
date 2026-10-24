"""New AT&T owners -> saved contacts, and a DM to Raf on their first sale.

WHY: Raf, #l10-alphalete 2026-10-22: "every time a new owner on AT&T Fiber or
AT&T NDS gets onboarded, I get their phone number saved in my contacts ... I get
a Slack message saying 'Hey, Shaylee Miller got a sale in her codes for the
first time' ... so I can text this owner 'Congratulations'". Eve, 2026-10-24,
split it in two:

  * ANY owner on Kelly's roster / in Nelrose's onboarding mails (the whole AT&T
    program, not just us) -> a contact in alphaletemarketing ONLY.
  * an owner of OUR org (on the Org workbook's "New Owners" tab, gid
    1674280247 — Eve's list) -> also a contact WITH PHONE in raffi127, and a DM
    from Lucy to Raf the day their codes log a first sale. Kelly's roster
    carries the whole program, so "on the roster" never means "ours" — the
    New Owners tab is the filter (Eve 2026-10-24: it was the Org Tree first).
    A row Eve marked Skip does not count.

WHERE THE DATA COMES FROM (both land in raffi127, read over IMAP like
fiber_owners_distro):
  Nelrose Lizada (nlizada@thesmartcircle.com), one mail per new owner:
      "AT&T NDS | Important Campaign Info (ICD Owner: X, ICD Company: Y)"
      -> name, company, campaign; the owner is the FIRST To address (the rest
      are their upline). NO phone in it.
  Nelrose's weekly "RES ATT Participants <dates>" (Fiber), one mail a week:
      a table ICD OWNER | ICD COMPANY | EMAIL | NUMBER -> name, email AND phone.
  Kelly Pavone's Friday "AT&T ICD Owner Roster" xlsx, every tab
      -> name, email, PHONE. So a phone can lag the onboarding by up to a week:
      the contact is saved without it and the phone is filled in the run after
      the roster that has it.

FIRST SALE = the first order under the owner's name in the order log:
  Fiber: ATT Tracker 2.1 - D2D > ORDERLOG > ALLREPS ('Owner Name')
  NDS:   NDS-SNRES-ATT-OOFWorkbook > ORDERLOG ('Owner & Office')
Only owners WATCHED from onboarding are judged, so "first" is honest: someone
picked up at setup who already has orders is marked as selling, never DM'd.

State: output/new_owner_contacts/state.json (one Lucy runs this).

  python -m automations.new_owner_contacts.run --dry-run   # read all, write nothing
  python -m automations.new_owner_contacts.run --apply     # the scheduled pass
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import email
import email.utils
import imaplib
import io
import json
import platform
import re
import sys
from html.parser import HTMLParser
from pathlib import Path
from typing import Dict, List, Optional, Tuple

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

REPO_ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = REPO_ROOT / "output" / "new_owner_contacts"
STATE_PATH = OUT_DIR / "state.json"

RAF_USER_ID = "U045Z8N0ZQC"           # Rafael Hidalgo
MARKETING_ACCOUNT = "alphaletemarketing@gmail.com"
RAF_ACCOUNT = "raffi127@gmail.com"

ONBOARD_SENDER = "nlizada@thesmartcircle.com"
ONBOARD_SUBJECT = "Important Campaign Info"
# Fiber has no per-owner mail: Nelrose sends ONE weekly "RES ATT Participants
# <dates>" to alphaletereporting + raffi127 + Maud, a table of that week's
# onboardings WITH phone (Eve 2026-10-24).
WEEKLY_SUBJECT = "Participants"
LOOKBACK_DAYS = 45      # how far back to read onboarding mails each pass
ORG_RECHECK_DAYS = 30   # an onboarding not on New Owners is re-checked this long (Eve may add the row late)
SALES_WINDOW_DAYS = 14  # order-log window read for first sales

_SUBJECT = re.compile(
    r"^\s*(?:\[[^\]]*\]\s*)?(?P<camp>[^|]+?)\s*\|\s*Important Campaign Info\s*"
    r"\(\s*ICD Owner:\s*(?P<owner>[^,]+?)\s*,\s*ICD Company:\s*(?P<co>.+?)\s*\)\s*$",
    re.I)


# --------------------------------------------------------------------------
# parsing (pure — tested)
# --------------------------------------------------------------------------

def norm_name(s: str) -> str:
    return " ".join(re.sub(r"[^a-z ]", " ", (s or "").lower()).split())


def parse_subject(subject: str) -> Optional[Dict[str, str]]:
    """Nelrose's onboarding subject -> {campaign, program, owner, company}.
    Replies/forwards are not onboardings (the original already counted)."""
    s = (subject or "").strip()
    if re.match(r"^(re|fw|fwd)\s*:", s, re.I):
        return None
    m = _SUBJECT.match(s)
    if not m:
        return None
    camp = " ".join(m.group("camp").split())
    return {"campaign": camp,
            "program": "NDS" if re.search(r"\bNDS\b", camp, re.I) else "Fiber",
            "owner": " ".join(m.group("owner").split()),
            "company": " ".join(m.group("co").split())}


class _Tables(HTMLParser):
    """Every <table> in an HTML mail as rows of cell text."""
    def __init__(self):
        super().__init__()
        self.tables: List[List[List[str]]] = []
        self._cell: Optional[List[str]] = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self.tables[-1].append([])
        elif tag in ("td", "th"):
            self._cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None:
            if self.tables and self.tables[-1]:
                self.tables[-1][-1].append(" ".join("".join(self._cell).split()))
            self._cell = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data.replace("\xa0", " "))


def parse_participants(html: str) -> List[Dict[str, str]]:
    """The weekly 'RES ATT Participants' table -> [{owner, company, email,
    phone}]. Columns found by header label (ICD OWNER / ICD COMPANY / EMAIL /
    NUMBER), never by position."""
    p = _Tables()
    p.feed(html or "")
    out = []
    for rows in p.tables:
        hdr_i = next((i for i, r in enumerate(rows)
                      if any(c.strip().upper() == "ICD OWNER" for c in r)), None)
        if hdr_i is None:
            continue
        col = {c.strip().upper(): j for j, c in enumerate(rows[hdr_i])}

        def g(r, label):
            j = col.get(label)
            return r[j].strip() if j is not None and j < len(r) else ""
        for r in rows[hdr_i + 1:]:
            owner = " ".join(g(r, "ICD OWNER").split())
            if not owner:
                continue
            out.append({"owner": owner, "company": g(r, "ICD COMPANY"),
                        "email": g(r, "EMAIL").lower(),
                        "phone": clean_phone(g(r, "NUMBER"))})
    return out


def clean_phone(raw) -> str:
    """Roster phone cell -> '(555) 123-4567' when it is a plain US number,
    else the cell as typed. openpyxl hands a numeric cell back as '5551234567.0'."""
    s = str(raw or "").strip()
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    digits = re.sub(r"\D", "", s)
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) == 10:
        return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
    return s


def split_name(full: str) -> Tuple[str, str]:
    parts = (full or "").split()
    if not parts:
        return "", ""
    return parts[0], " ".join(parts[1:])


def new_owner_names(rows) -> set:
    """Every name on Eve's New Owners tab, normalized — minus her Skip rows."""
    from automations.new_owners.bank import SKIP
    return {norm_name(r.name) for r in rows
            if r.name and r.status.strip().lower() != SKIP.lower()}


def tab_program(campaign: str) -> Optional[str]:
    """New Owners 'Campaign' -> the order log to watch. Only Fiber and NDS have
    one (Raf asked for those two); B2B / Retail JE rows -> None."""
    c = (campaign or "").lower()
    if "nds" in c:
        return "NDS"
    if "fiber" in c:
        return "Fiber"
    return None


def tab_watch_rows(rows) -> List[Tuple[str, str]]:
    """(name, program) for every Fiber/NDS row on the tab, minus Skip. These are
    watched even with NO Nelrose mail — Hasan Salame, Algemar Kennel and Vincent
    Dodson never had one (Eve 2026-10-24)."""
    from automations.new_owners.bank import SKIP
    out = []
    for r in rows:
        prog = tab_program(r.campaign)
        if r.name and prog and r.status.strip().lower() != SKIP.lower():
            out.append((" ".join(r.name.split()), prog))
    return out


def in_org(name: str, tree: set, aliases: Optional[dict] = None) -> bool:
    if norm_name(name) in tree:
        return True
    if aliases:
        from automations.focus_office_att.aliases import alias_to_canonical
        return norm_name(alias_to_canonical(name, aliases)) in tree
    return False


def _rows(text: str) -> List[List[str]]:
    first = text.split("\n", 1)[0]
    delim = "\t" if "\t" in first else ","
    return list(csv.reader(io.StringIO(text), delimiter=delim))


def _decode(raw: bytes) -> str:
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("latin-1")


def _parse_date(s: str) -> Optional[dt.date]:
    s = (s or "").strip().split(" ")[0]
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def first_orders(text: str, owner_headers=("Owner Name",),
                 owner_prefix: bool = False) -> Dict[str, dt.date]:
    """{normalized owner: earliest order date} from an order-log export.

    Columns by header label. `owner_prefix` is for NDS's 'Owner & Office'
    ('JANE DOE\\n[Tampa, FL]'): the owner is the text before the bracket
    (same cut as opt_nds._norm_owner)."""
    rows = _rows(text)
    if not rows:
        return {}
    head = [" ".join((h or "").lower().split()) for h in rows[0]]
    want = {" ".join(h.lower().split()) for h in owner_headers}
    oi = next((i for i, h in enumerate(head) if h in want), None)
    if oi is None and owner_prefix:
        oi = next((i for i, h in enumerate(head) if h.startswith("owner")), None)
    di = next((i for i, h in enumerate(head) if "order date" in h), None)
    if oi is None:
        raise ValueError(f"order log has no owner column: {rows[0][:6]}")
    out: Dict[str, dt.date] = {}
    for r in rows[1:]:
        if oi >= len(r):
            continue
        who = (r[oi] or "").strip()
        if owner_prefix:
            who = re.split(r"[\[\n]", who, maxsplit=1)[0]
        key = norm_name(who)
        if not key or key in ("total", "grand total"):
            continue
        d = _parse_date(r[di]) if di is not None and di < len(r) else None
        d = d or dt.date.today()
        if key not in out or d < out[key]:
            out[key] = d
    return out


def message(w: dict, sale_day: dt.date) -> str:
    """The DM to Raf. English — it is a Slack post (feedback_slack-channel-posts-in-english)."""
    lines = [f"Hey Raf, *{w['name']}* got a sale in their codes for the first "
             f"time ({w['program']}, {sale_day.month}/{sale_day.day})."]
    if w.get("phone"):
        lines.append(f"Phone: {w['phone']} (saved in your contacts)")
    else:
        lines.append("No phone on Kelly's roster yet.")
    if w.get("company"):
        lines.append(f"Company: {w['company']}")
    return "\n".join(lines)


# --------------------------------------------------------------------------
# sources
# --------------------------------------------------------------------------

def fetch_onboardings(days: int = LOOKBACK_DAYS) -> List[dict]:
    """Nelrose's onboarding mails from the last `days` days, oldest first."""
    from automations.fiber_owners_distro import roster as R
    since = (dt.date.today() - dt.timedelta(days=days)).strftime("%d-%b-%Y")
    M = imaplib.IMAP4_SSL(R.IMAP_HOST)
    M.login(R.ACCOUNT, R._app_password())
    out: List[dict] = []
    try:
        M.select('"[Gmail]/All Mail"', readonly=True)
        _typ, data = M.search(None, f'(FROM "{ONBOARD_SENDER}" '
                                    f'SUBJECT "{ONBOARD_SUBJECT}" SINCE {since})')
        for mid in (data[0] or b"").split():
            _t, md = M.fetch(mid, "(BODY.PEEK[HEADER.FIELDS (SUBJECT TO DATE)])")
            msg = email.message_from_bytes(md[0][1])
            info = parse_subject(R._decode(msg.get("Subject", "")))
            if not info:
                continue
            to = [a for _n, a in email.utils.getaddresses([msg.get("To", "")]) if a]
            info["email"] = (to[0].strip().lower() if to else "")
            try:
                info["date"] = email.utils.parsedate_to_datetime(
                    msg.get("Date")).date().isoformat()
            except Exception:  # noqa: BLE001
                info["date"] = dt.date.today().isoformat()
            out.append(info)
        _typ, data = M.search(None, f'(FROM "{ONBOARD_SENDER}" '
                                    f'SUBJECT "{WEEKLY_SUBJECT}" SINCE {since})')
        for mid in (data[0] or b"").split():
            _t, md = M.fetch(mid, "(BODY.PEEK[])")
            msg = email.message_from_bytes(md[0][1])
            subject = R._decode(msg.get("Subject", ""))
            if re.match(r"^(re|fw|fwd)\s*:", subject.strip(), re.I):
                continue
            html = next((part.get_payload(decode=True).decode(
                            part.get_content_charset() or "utf-8", "replace")
                         for part in msg.walk()
                         if part.get_content_type() == "text/html"), "")
            try:
                day = email.utils.parsedate_to_datetime(msg.get("Date")).date().isoformat()
            except Exception:  # noqa: BLE001
                day = dt.date.today().isoformat()
            nds = bool(re.search(r"\bNDS\b", subject, re.I))
            for row in parse_participants(html):
                out.append({**row, "date": day,
                            "campaign": "AT&T NDS" if nds else "RES ATT",
                            "program": "NDS" if nds else "Fiber"})
    finally:
        M.logout()
    return out


def roster_owners() -> Tuple[List, str]:
    """Every owner on every tab of Kelly's latest roster, + the roster's date."""
    import openpyxl
    from automations.fiber_owners_distro import roster as R
    path, edate = R.fetch_latest_roster(save_dir=OUT_DIR, verbose=False)
    wb = openpyxl.load_workbook(str(path), read_only=True)
    tabs = list(wb.sheetnames)
    wb.close()
    owners = []
    for tab in tabs:
        try:
            owners += R.parse_tab(path, tab)
        except RuntimeError as e:          # a tab with no Email column
            print(f"  roster tab {tab!r} skipped: {e}")
    return owners, edate.date().isoformat()


def read_new_owners() -> list:
    from automations.new_owners import bank as B
    _cmap, rows = B.read(B.open_bank(create=False))
    return rows


def pull_first_orders(programs: set, days: int = SALES_WINDOW_DAYS) -> Dict[str, Dict[str, dt.date]]:
    """{program: {owner: earliest order}} over the last `days` days."""
    end = dt.date.today()
    start = end - dt.timedelta(days=days)
    out: Dict[str, Dict[str, dt.date]] = {}
    if "Fiber" in programs:
        from automations.uploaded.order_log import ALLREPS_VIEW_URL_TMPL, CROSSTAB_SHEET
        from automations.shared.tableau_patchright import (
            download_crosstab_patchright, tableau_session)
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        tmp = OUT_DIR / "fiber_orderlog.csv"
        with tableau_session(headless=True, verbose=False) as page:
            download_crosstab_patchright(
                ALLREPS_VIEW_URL_TMPL.format(start=start.isoformat(),
                                             end=end.isoformat()),
                CROSSTAB_SHEET, tmp, verbose=False, page=page)
        out["Fiber"] = first_orders(_decode(tmp.read_bytes()), ("Owner Name",))
    if "NDS" in programs:
        from automations.office_metrics import nds_orderlog
        p = nds_orderlog.pull(days=days)
        out["NDS"] = first_orders(_decode(p.read_bytes()),
                                  ("Owner & Office", "Owner Name"), owner_prefix=True)
    for prog, m in out.items():
        print(f"  {prog} order log: {len(m)} owners with orders "
              f"({start}..{end})")
        if not m:
            raise ValueError(f"{prog} order log came back empty — not trusting it")
    return out


# --------------------------------------------------------------------------
# contacts
# --------------------------------------------------------------------------

def _contacts_index(svc) -> Dict[str, dict]:
    """canonical email -> {resourceName, etag, phones} for every contact."""
    from automations.fiber_owners_distro.contacts_write import _retry
    from automations.fiber_owners_distro.roster import canonical
    out, token = {}, None
    while True:
        resp = _retry(lambda: svc.people().connections().list(
            resourceName="people/me", pageSize=1000,
            personFields="emailAddresses,phoneNumbers",
            pageToken=token).execute())
        for p in resp.get("connections", []) or []:
            rec = {"resourceName": p.get("resourceName"), "etag": p.get("etag"),
                   "phones": [x.get("value") for x in p.get("phoneNumbers") or []]}
            for e in p.get("emailAddresses") or []:
                v = (e.get("value") or "").strip()
                if v:
                    out.setdefault(canonical(v), rec)
        token = resp.get("nextPageToken")
        if not token:
            return out


def plan_contacts(index: Dict[str, dict], people: List[dict]) -> Tuple[List[dict], List[dict]]:
    """(to create, to give a phone) — pure, tested. A contact that already has
    ANY phone is left alone: someone may have typed a better one by hand."""
    from automations.fiber_owners_distro.roster import canonical
    create, phone = [], []
    seen = set()
    for p in people:
        key = canonical(p["email"])
        if not key or key in seen:
            continue
        seen.add(key)
        rec = index.get(key)
        if rec is None:
            create.append(p)
        elif p.get("phone") and not rec["phones"]:
            phone.append({**p, **rec})
    return create, phone


def upsert_contacts(account: str, people: List[dict], *, apply: bool) -> dict:
    from automations.fiber_owners_distro.contacts_write import _retry, _service
    svc = _service(account)
    create, phone = plan_contacts(_contacts_index(svc), people)
    print(f"  {account}: {len(create)} new contact(s), {len(phone)} get a phone")
    for p in create[:15]:
        print(f"    + {p['name']} <{p['email']}> {p.get('phone') or ''}")
    if len(create) > 15:
        print(f"    ... and {len(create) - 15} more")
    for p in phone[:15]:
        print(f"    ☎ {p['name']} <{p['email']}> {p['phone']}")
    if not apply:
        return {"create": len(create), "phone": len(phone)}
    for i in range(0, len(create), 200):
        batch = []
        for p in create[i:i + 200]:
            given, family = split_name(p["name"])
            person = {"names": [{"givenName": given or p["email"],
                                 "familyName": family}],
                      "emailAddresses": [{"value": p["email"]}]}
            if p.get("company"):
                person["organizations"] = [{"name": p["company"]}]
            if p.get("phone"):
                person["phoneNumbers"] = [{"value": p["phone"], "type": "mobile"}]
            batch.append({"contactPerson": person})
        _retry(lambda: svc.people().batchCreateContacts(
            body={"contacts": batch, "readMask": "names"}).execute())
    for i in range(0, len(phone), 200):
        chunk = phone[i:i + 200]
        body = {"contacts": {p["resourceName"]: {
                    "etag": p["etag"],
                    "phoneNumbers": [{"value": p["phone"], "type": "mobile"}]}
                for p in chunk},
                "updateMask": "phoneNumbers", "readMask": "names"}
        _retry(lambda: svc.people().batchUpdateContacts(body=body).execute())
    return {"create": len(create), "phone": len(phone)}


# --------------------------------------------------------------------------
# state + Slack
# --------------------------------------------------------------------------

def load_state() -> dict:
    try:
        return json.loads(STATE_PATH.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001 — no file = first run
        return {}


def save_state(state: dict) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    STATE_PATH.write_text(json.dumps(state, indent=1, sort_keys=True),
                          encoding="utf-8")


def send(text: str) -> bool:
    """DM Raf as Lucy. Refuses from Windows / a non-Lucy token (atef_nds_watch)."""
    if platform.system() == "Windows":
        print("! not sending from Windows (token here is a person, not Lucy).")
        return False
    from automations.shared import slack_metrics_post as smp
    client = smp._client()
    who = ""
    try:
        who = client.auth_test().get("user", "")
    except Exception:  # noqa: BLE001
        pass
    if who and "lucy" not in who.lower():
        print(f"! this machine posts as {who}, not Lucy - not sending.")
        return False
    ch = client.conversations_open(users=RAF_USER_ID)["channel"]["id"]
    client.chat_postMessage(channel=ch, text=text)
    return True


# --------------------------------------------------------------------------
# the pass
# --------------------------------------------------------------------------

def run(*, apply: bool, skip_sales: bool = False) -> int:
    from automations.fiber_owners_distro.roster import canonical
    today = dt.date.today()
    state = load_state()
    first_run = not state
    watch: Dict[str, dict] = state.setdefault("watch", {})
    pending: Dict[str, dict] = state.setdefault("pending_org", {})

    onboard = fetch_onboardings()
    print(f"onboarding mails (Nelrose, last {LOOKBACK_DAYS}d): {len(onboard)}")
    owners, roster_date = roster_owners()
    print(f"Kelly's roster {roster_date}: {len(owners)} owner rows")
    phones, emails = {}, {}
    for o in owners:
        if o.email:
            emails.setdefault(norm_name(o.name), o.email.strip().lower())
        ph = clean_phone(o.phone)
        if ph:
            phones.setdefault(canonical(o.email), ph)
            phones.setdefault(norm_name(o.name), ph)

    def phone_for(email_: str, name: str, own: str = "") -> str:
        return own or phones.get(canonical(email_)) or phones.get(norm_name(name)) or ""

    # 1) everyone -> alphaletemarketing
    people = [{"name": o.name or o.email, "email": o.email, "company": o.company,
               "phone": clean_phone(o.phone)} for o in owners]
    people += [{"name": ob["owner"], "email": ob["email"], "company": ob["company"],
                "phone": phone_for(ob["email"], ob["owner"], ob.get("phone", ""))}
               for ob in onboard if ob["email"]]
    upsert_contacts(MARKETING_ACCOUNT, people, apply=apply)

    # 2) ours -> watch list (+ raffi127 below)
    tab_rows = read_new_owners()
    tree = new_owner_names(tab_rows)
    try:
        from automations.focus_office_att.aliases import load_aliases
        aliases = load_aliases()
    except Exception as e:  # noqa: BLE001 — the tree alone still works
        print(f"  (aliases unavailable: {e})")
        aliases = None
    print(f"New Owners tab: {len(tree)} names")
    def watched_as(name: str) -> Optional[str]:
        n = norm_name(name)
        return next((k for k, w in watch.items() if norm_name(w["name"]) == n), None)

    for ob in onboard:
        key = canonical(ob["email"]) or norm_name(ob["owner"])
        old = watched_as(ob["owner"])
        if key not in watch and old is not None:
            # Watched off the tab first; the mail brings the email -> re-key.
            watch[key] = watch.pop(old)
            watch[key]["email"] = watch[key].get("email") or ob["email"]
            watch[key]["company"] = watch[key].get("company") or ob["company"]
        if key in watch:
            watch[key]["phone"] = watch[key].get("phone") or phone_for(ob["email"], ob["owner"], ob.get("phone", ""))
            continue
        if in_org(ob["owner"], tree, aliases):
            watch[key] = {"name": ob["owner"], "email": ob["email"],
                          "company": ob["company"], "program": ob["program"],
                          "onboarded": ob["date"], "added": today.isoformat(),
                          "phone": phone_for(ob["email"], ob["owner"], ob.get("phone", "")),
                          "seeded": first_run}
            pending.pop(key, None)
            print(f"  watching {ob['owner']} ({ob['program']}, onboarded {ob['date']})")
        else:
            age = (today - dt.date.fromisoformat(ob["date"])).days
            if age <= ORG_RECHECK_DAYS:
                pending[key] = {"name": ob["owner"], "date": ob["date"]}
            else:
                pending.pop(key, None)
    # Every Fiber/NDS row on the tab is watched too, mail or no mail. Always
    # seeded: an owner already selling when the watch starts (Algemar was
    # active in September) must not read as a 'first' sale.
    for name, prog in tab_watch_rows(tab_rows):
        if watched_as(name) is not None:
            continue
        email_ = emails.get(norm_name(name), "")
        key = canonical(email_) if email_ else "name:" + norm_name(name)
        if key in watch:
            continue
        watch[key] = {"name": name, "email": email_, "company": "",
                      "program": prog, "onboarded": "", "added": today.isoformat(),
                      "phone": phone_for(email_, name), "seeded": True}
        print(f"  watching {name} ({prog}, from the New Owners tab)")
    if pending:
        print(f"  not ours (not on New Owners): "
              f"{', '.join(sorted(v['name'] for v in pending.values()))}")

    open_watch = {k: w for k, w in watch.items() if not w.get("first_sale")}
    raf_people = [{"name": w["name"], "email": w["email"], "company": w["company"],
                   "phone": w.get("phone", "")} for w in watch.values() if w.get("email")]
    if raf_people:
        upsert_contacts(RAF_ACCOUNT, raf_people, apply=apply)

    # 3) first sales -> DM Raf
    rc = 0
    if open_watch and not skip_sales:
        # A seeded owner is judged against the whole look-back, so an order
        # from before the watch existed can't later read as a 'first' sale.
        seeding = any(w.get("seeded") for w in open_watch.values())
        firsts = pull_first_orders({w["program"] for w in open_watch.values()},
                                   days=LOOKBACK_DAYS if seeding else SALES_WINDOW_DAYS)
        for key, w in sorted(open_watch.items()):
            got = firsts.get(w["program"], {})
            day = got.get(norm_name(w["name"]))
            if day is None and aliases:
                from automations.focus_office_att.aliases import get_search_candidates
                for cand in get_search_candidates(w["name"], aliases):
                    day = day or got.get(norm_name(cand))
            if day is None:
                print(f"  {w['name']}: no sale yet")
                continue
            w["first_sale"] = day.isoformat()
            if w.get("seeded"):
                w["dm"] = "skipped: already selling when the watch started"
                print(f"  {w['name']}: already selling at setup ({day}) - no DM")
                continue
            text = message(w, day)
            if not apply:
                print(f"\n--- would DM Raf ---\n{text}\n")
                continue
            if send(text):
                w["dm"] = today.isoformat()
                print(f"  DM'd Raf: {w['name']} first sale {day}")
            else:
                w.pop("first_sale", None)     # retry next pass
                rc = 1
        for w in open_watch.values():         # judged once: from now on a sale is news
            if not w.get("first_sale"):
                w["seeded"] = False
    elif not open_watch:
        print("nobody waiting on a first sale")

    if apply:
        save_state(state)
    return rc


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--apply", action="store_true",
                    help="write contacts, DM Raf, save state (the scheduled pass)")
    ap.add_argument("--dry-run", action="store_true",
                    help="read everything, print what would happen, write nothing")
    ap.add_argument("--skip-sales", action="store_true",
                    help="contacts only — don't pull the order logs")
    a = ap.parse_args(argv)
    return run(apply=a.apply and not a.dry_run, skip_sales=a.skip_sales)


if __name__ == "__main__":
    sys.exit(main())
