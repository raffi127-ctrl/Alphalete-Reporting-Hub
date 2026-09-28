# -*- coding: utf-8 -*-
"""Plain change-request for Carlos: what the ARS doc says now, what we
propose, and why. Megan 2026-09-27: "the revisions needed on the doc need to
be super simplified - telling carlos what it currently says and what we're
proposing/why"."""
import datetime as dt
from docx import Document
from docx.shared import Pt, RGBColor

GREY = RGBColor(0x5C, 0x65, 0x70)
RED = RGBColor(0xA8, 0x32, 0x2A)
GREEN = RGBColor(0x15, 0x6E, 0x46)

CHANGES = [
 ("Pay", "REBUTTALS table, lines 1035 and 1088",
  "We offer a weekly salary pay for the role (the base salary is determined "
  "on your background/experience) plus bonuses or commission.",
  "We offer weekly pay ranging from $1,000–$1,500 plus bonuses or "
  "commission. Is that something you're okay with?  (Or the paid-training "
  "answer on line 1036 — $16–$21 an hour plus commission. Either is fine; "
  "it is only the word \"base\" that has to go.)",
  "Raf: don't say there's a base. 25 messages in the last six weeks told "
  "applicants there is a base pay — 19 of them from one person — and the "
  "doc is where they got the wording."),

 ("Location", "REBUTTALS table, line 1038",
  "Our main office is located in *OfficeLocation*, however, you applied for "
  "our opening in *JobAdLocation*.",
  "You applied to our open [Role] position. Our office address is "
  "[the office's own address].",
  "Name the role first, then the address, in sentence case. One address per "
  "office and it does not change between rounds — Raf's candidates always "
  "go to Irving, Carlos's always to Grand Prairie."),

 ("Group interview", "REBUTTALS table line 1043, SMS Templates line 1080",
  "It looks like 2 other people also applied to the same listing as well. "
  "Depending on if they show up will determine if you would be conducting "
  "the interview by yourself or with someone else in the meeting.",
  "Yes, this will be a small group interview.",
  "Lead with yes and soften it. The current wording makes a group interview "
  "sound like an accident. The first round is always a group."),

 ("Company name", "throughout the doc",
  "ALPHALETE MARKETING, INC.",
  "Alphalete Marketing, Inc.",
  "All caps reads as shouting to an applicant. 212 messages in the last six "
  "weeks shout the company name, the address or the role — 78 from the AI, "
  "74 from recruiters, 60 from templates."),

 ("Second-interview address", "SMS Templates, 2nd Confirmation",
  "1901 N Highway 360 Suite 610, Grand Prairie, TX 75050 — presented as "
  "the address for second interviews.",
  "Label it as Carlos's office address (11580), used at every stage. "
  "Raf's offices use 3100 Premier Drive, Suite 207, Irving, Texas 75063 at "
  "every stage.",
  "It reads as a second-round venue for everyone. Six messages sent Raf's "
  "second-interview candidates to Grand Prairie."),
]


def build(path):
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Georgia"; st.font.size = Pt(11)
    r = d.add_paragraph().add_run("ARS Processes doc — proposed changes")
    r.bold = True; r.font.size = Pt(20)
    p = d.add_paragraph()
    r = p.add_run("For Carlos · {:%d %b %Y}".format(dt.date.today()))
    r.font.size = Pt(10); r.font.color.rgb = GREY

    # Carlos reads this cold. Say what prompted it and what the numbers
    # are before the first change, or the "why" lines have nothing to
    # stand on (Megan 2026-09-27: "so I can send to carlos").
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(
        "We pulled every applicant text across four recruiting accounts "
        "for the six weeks ending 25 September \u2014 about 55,000 messages "
        "\u2014 and checked what we actually reply against the approved "
        "answers in ARS Processes 2026. Several gaps trace back to the doc "
        "itself rather than to any one person \u2014 the pay line teaches "
        "wording we no longer use, and the second-interview address reads "
        "as though it applies to every office.")
    r.font.size = Pt(11)
    p = d.add_paragraph()
    r = p.add_run(
        "Five changes below. Each one gives the line as it stands, what we "
        "would put in its place, and the reason. Numbers are from those "
        "six weeks.")
    r.font.size = Pt(11)

    for n, (title, where, now, prop, why) in enumerate(CHANGES, 1):
        p = d.add_paragraph(); p.paragraph_format.space_before = Pt(18)
        r = p.add_run("{}. {}".format(n, title)); r.bold = True; r.font.size = Pt(14)
        p = d.add_paragraph()
        r = p.add_run(where); r.font.size = Pt(9); r.font.color.rgb = GREY

        for lbl, text, colour in (("Now:", now, RED),
                                  ("Change to:", prop, GREEN),
                                  ("Why:", why, None)):
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(lbl + " "); r.bold = True; r.font.size = Pt(10.5)
            if colour is not None:
                r.font.color.rgb = colour
            r2 = p.add_run(text); r2.font.size = Pt(10.5)
            if colour is not None:
                r2.font.color.rgb = colour
    d.save(path)
    return path


if __name__ == "__main__":
    print(build("output/ars-doc-proposed-changes.docx"))
