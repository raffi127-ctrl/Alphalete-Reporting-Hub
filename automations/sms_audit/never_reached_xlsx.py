# -*- coding: utf-8 -*-
"""Never-reached list as an .xlsx, one tab per office, to send to each team.

Megan 2026-09-28: "I need to send this to each team - put in an excel with
each office having it's own tab."

WHO IS ON IT. Everyone who had a text FAIL in the last two weeks and, across
the whole six weeks, has: no delivered text ever, no message from them ever,
and no booking. Anyone who replied, ever got a text through, or booked is
out — 1,564 people had a failure and only 213 pass all four tests.

WHAT IT CANNOT SEE, stated on every tab. The SMS log is texts. It has no
sight of phone calls or emails, and the First Interview Confirmation goes
out as Email/SMS, so some of these people were emailed and some were rung
(LM1/LM2/LM3 is exactly that). So the honest claim is "no evidence of
contact in the text log", never "never contacted", and the tab says so
above the table rather than leaving a team to assume.

Records with NO phone number are left off entirely (Megan 2026-09-28: "you
can leave off the ones with no phone numbers"). A record with phone
0000000000 retried seventeen times is a data-entry fix, not a lost
applicant, and a team cannot act on it. The count dropped is printed and
shown on the tab, so they vanish from the list without vanishing from the
record.

  python -m automations.sms_audit.never_reached_xlsx
"""
from __future__ import annotations

import collections
import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from automations.sms_audit import analyze as A

OUTPUT = Path(__file__).resolve().parents[2] / "output"
ALL_WEEKS = ["w0821", "w0828", "w0904", "w0911", "w0918", "w0925"]
RECENT = ["w0918", "w0925"]
OFFICES = collections.OrderedDict([
    ("11280", "Rafael Hidalgo 11280"),
    ("23965", "Rafael 2nd Funnel 23965"),
    ("24065", "Raf New Recruiter 24065"),
    ("11580", "Carlos Hidalgo 11580"),
])
HEADERS = ["Applicant", "Phone", "Call List Status", "Texts Tried",
           "Why They Failed", "First Try", "Last Try", "Note"]
WIDTHS = [26, 15, 20, 12, 28, 16, 16, 26]

# The Call Hub's own export, dropped in Downloads. It is an HTML table
# named .xls, one file per office: callList_<office>.xls. Scraping the hub
# got 40 rows of 1,307 — the export gives all of them, so it is the source.
EXPORT_DIRS = [Path.home() / "Downloads", OUTPUT]


def call_status():
    """{office: {last-10-digits: status}} from whatever exports are on disk."""
    import html
    import re
    out = {}
    for folder in EXPORT_DIRS:
        if not folder.exists():
            continue
        for f in sorted(folder.glob("callList_*.xls")):
            office = f.stem.split("_")[-1]
            text = f.read_text(encoding="utf-8", errors="replace")
            rows = re.findall(r"<tr[^>]*>(.*?)</tr>", text, re.S | re.I)
            book = out.setdefault(office, {})
            for row in rows[1:]:
                cells = [html.unescape(re.sub(r"<[^>]+>", "", c)).strip()
                         for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>",
                                             row, re.S | re.I)]
                if len(cells) < 10:
                    continue
                digits = "".join(ch for ch in cells[4] if ch.isdigit())[-10:]
                if digits:
                    book[digits] = cells[-1]
    return out

INK = Font(name="Georgia", size=12, bold=True, color="000000")
HEAD = Font(name="Georgia", size=12, bold=True, color="FFFFFF")
NOTE = Font(name="Georgia", size=10, bold=False, italic=True, color="555555")
FILL = PatternFill("solid", fgColor="2F5D57")
DIM = PatternFill("solid", fgColor="F2F0EA")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
EDGE = Side(style="thin", color="BFBDB6")
BOX = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)

BAD_PHONE = ("0000000000", "", "1111111111")


def collect(status_by_office=None):
    """{office: [rows]} — everyone with no evidence of contact."""
    out = collections.OrderedDict((o, []) for o in OFFICES)
    dropped = collections.Counter()
    status_by_office = status_by_office or {}
    for office in OFFICES:
        per = collections.defaultdict(
            lambda: {"fail": 0, "ok": 0, "in": 0, "booked": False, "name": "",
                     "att": 0, "why": collections.Counter(),
                     "first": None, "last": None})
        for tag in ALL_WEEKS:
            recs, _s = A.load_office(office, tag)
            log, _s2 = A.load_log(office, tag)
            if not recs or not log:
                continue
            for phone, c in A.log_conversations(
                    log, A.booked_index(recs)).items():
                d = per[phone]
                d["name"] = d["name"] or (c.get("name") or "")
                if c["booked"]:
                    d["booked"] = True
                for m in c["msgs"]:
                    if m["dir"] == "In":
                        d["in"] += 1
                        continue
                    d["att"] += 1
                    status = (m.get("status") or "").strip()
                    if status.lower() == "delivered":
                        d["ok"] += 1
                    else:
                        if tag in RECENT:
                            d["fail"] += 1
                        d["why"][status or "(blank)"] += 1
                    when = m["when"]
                    if d["first"] is None or when < d["first"]:
                        d["first"] = when
                    if d["last"] is None or when > d["last"]:
                        d["last"] = when
        for phone, d in per.items():
            if not (d["fail"] and d["ok"] == 0 and d["in"] == 0
                    and not d["booked"]):
                continue
            # No usable number at all — not a person a team can chase.
            digits = "".join(ch for ch in (phone or "") if ch.isdigit())
            if (phone in BAD_PHONE or len(digits) < 10
                    or len(set(digits)) <= 1
                    or "Dummy" in " ".join(d["why"])):
                dropped[office] += 1
                continue
            dummy = "Not Valid" in " ".join(d["why"])
            book = status_by_office.get(office) or {}
            out[office].append({
                "status": book.get(phone[-10:], "" if book
                                   else "(no export loaded)"),
                "name": (d["name"] or "(no name on file)").title(),
                "phone": phone or "(none)",
                "att": d["att"],
                "why": "; ".join("{} x{}".format(k, v)
                                 for k, v in d["why"].most_common()),
                "first": d["first"], "last": d["last"],
                "note": ("Number rejected as invalid" if dummy else ""),
                "dummy": dummy})
        # real people first; a recruiter should not wade through dead records
        out[office].sort(key=lambda r: (r["dummy"], -r["att"], r["name"]))
    return out, dropped


def build(path=None):
    statuses = call_status()
    data, dropped = collect(statuses)
    wb = Workbook()
    wb.remove(wb.active)
    for office, label in OFFICES.items():
        rows = data[office]
        ws = wb.create_sheet(label[:31])
        ws["A1"] = "Applicants We Could Not Reach — {}".format(label)
        ws["A1"].font = Font(name="Georgia", size=14, bold=True)
        ws["A2"] = (
            "Every text to these people failed in the two weeks to 25 Sep, "
            "they never replied, and they never booked. "
            "This is the TEXT log only — it cannot see phone calls or "
            "emails, so treat it as “no evidence of contact by text” "
            "rather than “never contacted”.")
        ws["A2"].font = NOTE
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=7)
        line = "{} {}".format(len(rows),
                              "person" if len(rows) == 1 else "people")
        if dropped[office]:
            line += "   (plus {} record{} with no usable phone number, left "
            line = line.format(dropped[office],
                               "" if dropped[office] == 1 else "s")
            line += "off — those are records to fix, not people to chase)"
        ws["A5"] = line
        ws["A5"].font = Font(name="Georgia", size=12, bold=True)

        for i, h in enumerate(HEADERS, 1):
            cell = ws.cell(row=6, column=i, value=h)
            cell.font = HEAD
            cell.fill = FILL
            cell.alignment = CENTER
            cell.border = BOX
            ws.column_dimensions[get_column_letter(i)].width = WIDTHS[i - 1]
        for r, row in enumerate(rows, start=7):
            vals = [row["name"], row["phone"],
                    row["status"] or "not on the call list", row["att"],
                    row["why"],
                    row["first"].strftime("%d %b %H:%M") if row["first"] else "",
                    row["last"].strftime("%d %b %H:%M") if row["last"] else "",
                    row["note"]]
            for i, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=i, value=v)
                cell.font = INK
                cell.alignment = CENTER
                cell.border = BOX
                if row["dummy"]:
                    cell.fill = DIM
        ws.freeze_panes = "A7"
        ws.row_dimensions[2].height = 30
    path = path or (OUTPUT / "applicants-we-could-not-reach.xlsx")
    OUTPUT.mkdir(exist_ok=True)
    wb.save(str(path))
    return path, {o: len(v) for o, v in data.items()}, dropped


if __name__ == "__main__":
    p, counts, dropped = build()
    print("wrote {}".format(p))
    for o, n in counts.items():
        print("   {:<26} {:>4}   (dropped {} with no number)".format(
            OFFICES[o], n, dropped.get(o, 0)))
    sys.exit(0)
