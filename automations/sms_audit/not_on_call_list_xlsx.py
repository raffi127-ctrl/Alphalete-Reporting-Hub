# -*- coding: utf-8 -*-
"""Two workbooks — the people no text reached AND who are not on the call
list at all, one file per office so each can be sent on its own.

Megan 2026-09-28: "so I need those listed out in 2 excel sheets (1 for each)".
Only 11280 and 11580 have any; 23965 and 24065 called everybody.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from automations.sms_audit.never_reached_xlsx import (
    OFFICES, OUTPUT, call_status, collect)

INK = Font(name="Georgia", size=12, bold=True, color="000000")
HEAD = Font(name="Georgia", size=12, bold=True, color="FFFFFF")
NOTE = Font(name="Georgia", size=10, italic=True, color="555555")
FILL = PatternFill("solid", fgColor="8C3B2E")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
E = Side(style="thin", color="BFBDB6")
BOX = Border(left=E, right=E, top=E, bottom=E)
COLS = ["Applicant", "Phone", "Texts Tried", "Why They Failed",
        "First Try", "Last Try"]
W = [26, 16, 12, 30, 18, 18]

statuses = call_status()
data, _dropped = collect(statuses)
made = []
for office, label in OFFICES.items():
    rows = [r for r in data[office]
            if (r["status"] or "").strip() in ("", "not on the call list")]
    if not rows:
        print("{:<26} none".format(label))
        continue
    wb = Workbook()
    ws = wb.active
    ws.title = "Not On Call List"
    ws["A1"] = "Not Reached And Not On The Call List — {}".format(label)
    ws["A1"].font = Font(name="Georgia", size=14, bold=True)
    ws["A2"] = ("Every text to these people failed, they never replied, they "
                "never booked, AND they are not in the call list — so nobody "
                "is going to ring them either. Texts and the call list are "
                "all this can see; it cannot see email.")
    ws["A2"].font = NOTE
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=6)
    ws["A5"] = "{} {}".format(len(rows),
                              "person" if len(rows) == 1 else "people")
    ws["A5"].font = Font(name="Georgia", size=12, bold=True)
    for i, h in enumerate(COLS, 1):
        c = ws.cell(row=6, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = HEAD, FILL, CENTER, BOX
        ws.column_dimensions[get_column_letter(i)].width = W[i - 1]
    for r, row in enumerate(rows, start=7):
        vals = [row["name"], row["phone"], row["att"], row["why"],
                row["first"].strftime("%d %b %H:%M") if row["first"] else "",
                row["last"].strftime("%d %b %H:%M") if row["last"] else ""]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.alignment, c.border = INK, CENTER, BOX
    ws.freeze_panes = "A7"
    ws.row_dimensions[2].height = 30
    path = OUTPUT / "not-on-call-list-{}.xlsx".format(office)
    wb.save(str(path))
    made.append(str(path))
    print("{:<26} {:>3} -> {}".format(label, len(rows), path.name))
