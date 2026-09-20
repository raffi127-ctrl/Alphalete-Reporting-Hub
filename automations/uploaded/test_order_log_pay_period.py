"""Pay Period calendar on the Order Log (Raf 2026-09-16).

  python -m unittest automations.uploaded.test_order_log_pay_period
"""
import unittest
from datetime import date

import pandas as pd
from openpyxl import Workbook

from automations.uploaded import order_log as ol


class PayPeriodWeeks(unittest.TestCase):
    def test_window_sunday_to_saturday_paid_next_friday(self):
        weeks = ol._pay_period_weeks(date(2026, 9, 16))       # a Wednesday
        self.assertEqual(len(weeks), ol.PAY_PERIOD_WEEKS)
        self.assertEqual(weeks[0], (date(2026, 8, 23), date(2026, 8, 29),
                                    date(2026, 9, 4)))
        self.assertIn((date(2026, 9, 13), date(2026, 9, 19),
                       date(2026, 9, 25)), weeks)
        for sun, sat, fri in weeks:
            self.assertEqual((sun.weekday(), sat.weekday(), fri.weekday()),
                             (6, 5, 4))

    def test_tab_is_the_first_one(self):
        """Raf 2026-09-19: Pay Period first, and the file opens on it."""
        wb = Workbook()
        wb.active.title = "Cleaned Order Log"
        wb.create_sheet("Aaron Corona")
        ol._add_pay_period_tab(wb, date(2026, 9, 16))
        self.assertEqual(wb.sheetnames[0], ol.PAY_PERIOD_TAB)
        self.assertEqual(wb.active.title, ol.PAY_PERIOD_TAB)
        sh = wb[ol.PAY_PERIOD_TAB]
        self.assertEqual(sh["A1"].value, "Activation/Pay Week")
        self.assertEqual(sh["A4"].value, date(2026, 8, 23))


TODAY = date(2026, 9, 19)


def _log(rows):
    """A cleaned-log frame: (rep, order date, status[, install date])."""
    df = pd.DataFrame([{h: None for h in ol.FRIENDLY_HEADERS} for _ in rows])
    for i, row in enumerate(rows):
        rep, d, status = row[:3]
        df.at[i, "Rep"] = rep
        df.at[i, "Order Date"] = d
        df.at[i, "Status"] = status
        df.at[i, "Install Date"] = row[3] if len(row) > 3 else None
        df.at[i, "Product Type"] = "NEW INTERNET"
        df.at[i, "Package"] = "1 GIG"
    return df


class QuietRepsGetNoTab(unittest.TestCase):
    """Raf 2026-09-19: no tab for a rep who hasn't sold in two weeks."""

    def _tabs(self, rows, today=TODAY):
        wb = Workbook()
        wb.active.title = "Cleaned Order Log"
        ol._append_rep_breakdown_tabs(wb, _log(rows), today=today)
        return wb.sheetnames[1:]

    def test_cutoff_is_two_sunday_weeks_back(self):
        # Newest sale 9/16 (Wed) -> keep from Sunday 9/6 on.
        self.assertEqual(ol._rep_tab_cutoff(date(2026, 9, 16)), date(2026, 9, 6))

    def test_sold_this_week_and_last_week_keep_their_tabs(self):
        tabs = self._tabs([
            ("Selling Sam", date(2026, 9, 16), "Active"),      # this week
            ("Last Week Lou", date(2026, 9, 9), "Active"),     # week before
            ("Quiet Quinn", date(2026, 9, 5), "Active"),       # one day too old
            ("Gone Gabe", date(2026, 8, 3), "Active"),
        ])
        self.assertEqual(tabs, ["Selling Sam", "Last Week Lou"])

    def test_a_pending_line_still_counts_as_selling(self):
        tabs = self._tabs([
            ("Selling Sam", date(2026, 9, 16), "Active"),
            ("Pending Pete", date(2026, 9, 14), "Pending"),
        ])
        self.assertIn("Pending Pete", tabs)

    def test_a_scheduled_install_doesnt_move_the_cutoff(self):
        """Installs booked weeks out must not drop everyone who sold today."""
        tabs = self._tabs([
            ("Selling Sam", date(2026, 9, 16), "Active", date(2026, 10, 14)),
            ("Last Week Lou", date(2026, 9, 9), "Active"),
        ])
        self.assertEqual(tabs, ["Selling Sam", "Last Week Lou"])

    def test_cutoff_follows_the_file_not_today(self):
        """An old export still writes the tabs it would have written then."""
        tabs = self._tabs([
            ("Old Olivia", date(2025, 3, 5), "Active"),
            ("Older Oscar", date(2025, 1, 8), "Active"),
        ])
        self.assertEqual(tabs, ["Old Olivia"])


if __name__ == "__main__":
    unittest.main()
