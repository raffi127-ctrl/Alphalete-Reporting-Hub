"""The NDS Order Log workbook, B2B-style (Carlos 2026-10-29: "build it
that way. also include the customers first and last name in it").

Four kinds of tabs, same shape as the AT&T B2B log:
  1. All Reps        — one row per LINE from the wireless per-line feed,
                       customer first+last on every row
  2. Pending by Rep  — lines still in flight (not Active, not dead)
  3. Paycheck by Week— reps down, DD weeks across, PAID $ straight from
                       the deposit detail (Scoutly RAW, NDS Wireless) —
                       the actual paychecks, not an estimate
  4. one tab per rep — that rep's lines, Rep column dropped
"""
from __future__ import annotations

import collections
import datetime as dt
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Same palette + status logic as the Alphalete sales ATT log (Carlos
# 2026-10-29: "we should make it like that one is made") — att_order_log
# colors.py is the shared authority, so a status tweak lands in both.
HEADER_BG, WEEK_BG = "434343", "2563EB"
DEAD = {"disconnected", "cancelled", "canceled", "returned"}
HDR = ["Rep", "Customer", "Order Date", "DSI #", "Wireless Order #",
       "BAN", "Line", "Type", "Device", "Plan", "Line Status",
       "Ship Date", "Delivered", "Port Status"]


def _fill_for(status: str):
    from automations.att_order_log import colors as attc
    s = (status or "").strip().lower()
    if s == "active":
        return "C6EFCE"                       # his Posted green
    hexfill = attc.fill_for(status or "")
    if hexfill:
        return hexfill.lstrip("#")
    return "FFF2CC" if s else None            # unmapped in-flight = yellow


def _nice(name: str) -> str:
    from automations.shared.name_case import titlecase_name
    return titlecase_name((name or "").strip())


def _row_of(l: dict):
    return [_nice(l.get("rep", "")), _nice(l.get("customer", "")),
            l.get("order_date", ""), l.get("dsi", ""),
            l.get("wireless_order", ""), l.get("ban", ""),
            l.get("line_no", ""), l.get("line_type", ""),
            l.get("device", ""), l.get("plan", ""),
            l.get("line_status", ""), l.get("ship_date", ""),
            l.get("delivered", ""), l.get("port_status", "")]


def _thin():
    from openpyxl.styles import Border, Side
    side = Side(style="thin", color="D9D9D9")
    return Border(left=side, right=side, top=side, bottom=side)


def _header_row(sh, row, labels):
    for ci, label in enumerate(labels, start=1):
        c = sh.cell(row=row, column=ci, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADER_BG)
        c.border = _thin()


def _banner(sh, row, text, span):
    from openpyxl.utils import get_column_letter as gcl
    sh.merge_cells(f"A{row}:{gcl(span)}{row}")
    c = sh.cell(row=row, column=1, value=text)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=WEEK_BG)


def _data_row(sh, row, l, labels, drop_rep):
    vals = _row_of(l)
    if drop_rep:
        vals = vals[1:]
    fill = _fill_for(l.get("line_status", ""))
    for ci, v in enumerate(vals[:len(labels)], start=1):
        c = sh.cell(row=row, column=ci, value=v)
        c.border = _thin()
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)


def _act_week(l):
    """Saturday-start week of the line's delivered date (the activation
    proxy until the daily snapshots carry the real flip day); None =
    not yet activated."""
    if (l.get("line_status") or "").strip().lower() != "active":
        return None
    d = l.get("delivered") or ""
    try:
        dd = dt.datetime.strptime(d, "%m/%d/%Y").date()
    except ValueError:
        return None
    return dd - dt.timedelta(days=(dd.weekday() - 5) % 7)


def _autosize(sh, labels):
    for ci in range(1, len(labels) + 1):
        w = max([len(str(labels[ci - 1]))] +
                [len(str(sh.cell(row=r, column=ci).value or ""))
                 for r in range(1, sh.max_row + 1)
                 if not isinstance(sh.cell(row=r, column=ci),
                                   type(None))])
        sh.column_dimensions[get_column_letter(ci)].width = min(w + 2, 42)


def _write_flat(sh, lines, drop_rep=False):
    labels = HDR[1:] if drop_rep else HDR
    _header_row(sh, 1, labels)
    r = 2
    for l in lines:
        _data_row(sh, r, l, labels, drop_rep)
        r += 1
    _autosize(sh, labels)
    sh.freeze_panes = "A2"


def _write_rep_tab(sh, lines):
    """The ATT-log layout: a blue banner per activated week (newest
    first), then Not Yet Activated, dark headers under each banner."""
    labels = HDR[1:]
    groups = collections.defaultdict(list)
    for l in lines:
        groups[_act_week(l)].append(l)
    weeks = sorted((w for w in groups if w), reverse=True)
    row = 1
    for wk in weeks + ([None] if None in groups else []):
        g = groups[wk]
        if wk:
            text = (f"Activated Week {wk.strftime('%m/%d')} - "
                    f"{(wk + dt.timedelta(days=6)).strftime('%m/%d/%Y')}"
                    f"  •  {len(g)} line{'' if len(g) == 1 else 's'}")
        else:
            text = (f"Not Yet Activated  •  {len(g)} "
                    f"line{'' if len(g) == 1 else 's'}")
        _banner(sh, row, text, len(labels))
        row += 1
        _header_row(sh, row, labels)
        row += 1
        for l in g:
            _data_row(sh, row, l, labels, drop_rep=True)
            row += 1
        row += 2
    _autosize(sh, labels)
    sh.freeze_panes = "A3"


def _safe_title(name, used):
    t = re.sub(r"[\\/?*\[\]:]", " ", str(name)).strip()[:28] or "Sheet"
    base, n = t, 2
    while t.lower() in used:
        t = f"{base[:25]} {n}"
        n += 1
    used.add(t.lower())
    return t


def _paycheck_tab(wb, used, log):
    """Reps x DD-week $ matrix from the Scoutly RAW tab — the real
    paychecks. Fails soft: no sheet access leaves the tab out, never
    sinks the workbook."""
    try:
        from automations.recruiting_report.fill import open_by_key
        sh = open_by_key("1Hltk25zTudsaoYJFKvKqWlpT_4MF5_ZZq734XKVCJKY")
        raw = next(w for w in sh.worksheets()
                   if w.title.strip().lower() == "raw")
        vals = raw.get_all_values()
    except Exception as e:  # noqa: BLE001
        log(f"[ndsx] paycheck tab skipped ({type(e).__name__})")
        return
    hix = {h: i for i, h in enumerate(vals[0])}

    def money(v):
        try:
            return float(str(v).replace("$", "").replace(",", "") or 0)
        except ValueError:
            return 0.0

    paid = collections.defaultdict(float)
    weeks, reps = [], set()
    for r in vals[1:]:
        if not (r and r[0].strip()
                and r[hix["Campaign"]].strip() == "NDS Wireless"):
            continue
        wk, rep = r[0].strip(), _nice(r[hix["Rep Name"]])
        if wk not in weeks:
            weeks.append(wk)
        reps.add(rep)
        paid[(rep, wk)] += money(r[hix["Total $ to ICD"]])
    if not reps:
        return
    psh = wb.create_sheet(_safe_title("Paycheck by Week", used))
    psh.append(["Paid $ per rep per DD week (straight from the deposit "
                "detail — NDS Wireless, all buckets)"])
    psh.cell(row=1, column=1).font = Font(bold=True)
    psh.append(["Rep"] + [f"WE {w}" for w in weeks])
    for c in psh[2]:
        c.font = Font(bold=True)
    for rep in sorted(reps, key=lambda x: -sum(
            paid.get((x, w), 0) for w in weeks)):
        psh.append([rep] + [round(paid.get((rep, w), 0)) or ""
                            for w in weeks])
    psh.append(["TOTAL"] + [round(sum(paid.get((r, w), 0) for r in reps))
                            for w in weeks])
    for c in psh[psh.max_row]:
        c.font = Font(bold=True)
    psh.freeze_panes = "B3"


def build(orders, lines, today: dt.date, out_dir: Path, log=print) -> Path:
    out = out_dir / f"NDS Order Log {today.strftime('%m-%d-%Y')}.xlsx"
    wb = Workbook()
    used: set = set()
    lines = sorted(lines, key=lambda l: (_nice(l.get("rep", "")),
                                         l.get("order_date", ""),
                                         l.get("dsi", "")))
    # 1. All Reps — flat, sorted by rep (Megan's rule on the ATT log)
    sh = wb.active
    sh.title = _safe_title("All Reps", used)
    _write_flat(sh, lines)
    # 2. Pending by Rep — in-flight only
    pend = [l for l in lines
            if (l.get("line_status") or "").lower() not in
            ({"active"} | DEAD)]
    _write_flat(wb.create_sheet(_safe_title("Pending by Rep", used)), pend)
    # 3. Paycheck by Week (from the DD)
    _paycheck_tab(wb, used, log)
    # 4. one tab per rep, week-banner layout like the ATT log
    by_rep = collections.defaultdict(list)
    for l in lines:
        r = _nice(l.get("rep", ""))
        if r:
            by_rep[r].append(l)
    for rep in sorted(by_rep):
        _write_rep_tab(wb.create_sheet(_safe_title(rep, used)),
                       by_rep[rep])
    wb.save(out)
    log(f"[ndsx] order log workbook: {len(lines)} line(s), "
        f"{len(by_rep)} rep tab(s), pending {len(pend)}")
    return out
