"""Fiber Metrics — Carlos's crew on Raf's codes, from the D2D order log.

Carlos 2026-10-08: he sells under Raf's codes now ("Fiber"), not B2B logins —
"the majority of these reports remade ... specifically just for my reps."
Phase 1 = the boards that need no comp knowledge, all from ONE 60-day pull of
ATTTRACKER2_1-D2D 'A.Order Log' (the same crosstab the commission sheet
pulls), filtered to Rafael Hidalgo's office and then to Carlos's crew (the
GUEST_REPS list that already drives FIB-1/3 + SAL-7/8):

  order_log        the crew's lines as a status-colored .xlsx
  activation_by_rep  activated/sold by sale-date age (0-30 / 31-60), crew
                   total first, B2B color bands — the same board just built
                   for Khalil's NDS office
  pending_orders   delivered-nothing-yet chase list (Confirmed / Shipped /
                   blank status, no activation), Box-style per-rep bands
  customer_churn   the 0-30 rolloff list from the log's own Churn Buckets /
                   Days to Disconnect columns
  churn_rates      crew churn per product per bucket (computed — the D2D
                   CHURN views are office-level, not crew-level)

Fiber-log facts (probed 2026-10-08): statuses Confirmed / Posted / Shipped /
Disconnected / Canceled / blank; Posted = activated; spe.Name (SPE-########)
is the DD join key (cl.Production Lookup, ~96%% per commission_sheet); native
Churn Buckets, Days to Disconnect, Disconnect Reason, Bonus Eligible $ cols.

Destination (Carlos 2026-10-08): its own daily "Fiber Metrics — {date}"
thread in #alphalete-gp-sales — title-only parent, contents as the first
reply, every post starting with the campaign name per the callouts README.
NOTHING posts to the channel until Carlos approves the DM previews.

  lucy rerun fiber_metrics -- --dm U046G04P5LG        # previews, no channel
  lucy rerun fiber_metrics -- --post                  # the live thread
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import json
import sys
from pathlib import Path
from typing import Dict, List, Optional

REPO_ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = REPO_ROOT / "output" / "fiber_metrics"
CHANNEL = ("#alphalete-gp-sales", "C07J46MQNUX")
CARLOS = "U046G04P5LG"
LOOKBACK_DAYS = 60
BUCKETS = ("0-30 Day", "31-60 Day")
_BAND_KEY = {"0-30 Day": "0-30", "31-60 Day": "31-60"}
POSTED = ("posted",)
DEAD = ("canceled", "cancelled", "disconnected")


# ---------------------------------------------------------------- load ----

def _parse_date(v) -> Optional[dt.date]:
    s = str(v or "").strip()
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def pull(today: dt.date, log=print) -> Path:
    """The D2D A.Order Log crosstab with OUR 60-day window."""
    from automations.commission_sheet.sources import read_crosstab  # noqa: F401
    from automations.harvest import adapter as _hv
    from automations.uploaded import order_log as _ol
    start = today - dt.timedelta(days=LOOKBACK_DAYS)
    out = OUT_DIR / f"d2d_orderlog_{today.isoformat()}.csv"
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists() and out.stat().st_size > 1000:
        log(f"[fiber] reusing today's pull ({out.stat().st_size:,} bytes)")
        return out
    url = _ol.ALLREPS_VIEW_URL_TMPL.format(start=start.isoformat(),
                                           end=today.isoformat())
    if _hv.try_cache_view(url, _ol.CROSSTAB_SHEET, out) is None:
        from automations.shared.tableau_patchright import (
            download_crosstab_patchright)
        download_crosstab_patchright(url, _ol.CROSSTAB_SHEET, out)
    return out


class Cols:
    """Tolerant column lookup — the crosstab reshapes; names, not indices."""

    def __init__(self, header: List[str]):
        self.header = [str(h or "").strip() for h in header]
        self._norm = [" ".join(h.lower().split()) for h in self.header]

    def find(self, *cands) -> Optional[int]:
        for c in cands:
            c = " ".join(c.lower().split())
            if c in self._norm:
                return self._norm.index(c)
        for c in cands:
            c = " ".join(c.lower().split())
            for i, h in enumerate(self._norm):
                if c in h:
                    return i
        return None

    def need(self, label, *cands) -> int:
        i = self.find(*cands)
        if i is None:
            raise RuntimeError(f"fiber log: no {label} column in "
                               f"{self.header}")
        return i


def load_crew(path: Path, log=print):
    """-> (Cols, crew rows). Rafael's office, then the GUEST_REPS crew."""
    from automations.commission_sheet.sources import read_crosstab
    from automations.total_knocks import guests
    rows = read_crosstab(path)
    hdr_i = next(i for i, r in enumerate(rows)
                 if any(str(c).strip().lower() == "rep" for c in r))
    cols = Cols(rows[hdr_i])
    data = rows[hdr_i + 1:]
    i_owner = cols.need("owner", "owner name", "owner & office", "owner")
    i_rep = cols.need("rep", "rep")
    raf = [r for r in data if len(r) > i_owner
           and "RAFAEL" in str(r[i_owner]).upper()
           and "HIDALGO" in str(r[i_owner]).upper()]
    crew_names = guests.roster("Rafael Hidalgo", "Carlos Hidalgo")
    crew_toks = [guests._tokens(n) for n in crew_names]

    def _is_crew(name: str) -> bool:
        t = guests._tokens(name)
        return any(guests._subseq(ct, t) or guests._subseq(t, ct)
                   for ct in crew_toks)

    crew = [r for r in raf
            if len(r) > i_rep and _is_crew(str(r[i_rep]).strip())]
    log(f"[fiber] {len(data)} rows -> rafael {len(raf)} -> crew {len(crew)} "
        f"({len(crew_names)} names)")
    if not crew:
        raise RuntimeError("fiber log: ZERO crew rows — check the pull "
                           "window / crew spellings before posting anything")
    return cols, crew


def _c(cols: Cols, r, i: Optional[int]) -> str:
    return (str(r[i]).strip() if i is not None and len(r) > i and r[i]
            else "")


# ---------------------------------------------------------------- boards --

def build_order_log(cols, crew, today, log=print) -> Path:
    """The crew's lines as a status-colored workbook."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    FILL = {"posted": "C6EFCE", "confirmed": "FFF2CC", "shipped": "FFF2CC",
            "disconnected": "FFC7CE", "canceled": "FFC7CE",
            "cancelled": "FFC7CE"}
    want = [
        ("Rep", ("rep",)),
        ("Order Date", ("sp.order date",)),
        ("Customer", ("customer name",)),
        ("SPM", ("sp.spm number",)),
        ("SPE", ("spe.name",)),
        ("Status", ("dtr status", "spe.status")),
        ("First Available", ("dtr first available",)),
        ("Product", ("product type (broken out)", "product")),
        ("Provider", ("spe.provider",)),
        ("Churn Bucket", ("churn buckets",)),
        ("Days to Disc", ("days to disconnect",)),
        ("Disc Reason", ("disconnect reason",)),
        ("Bonus Eligible $", ("bonus eligible",)),
    ]
    ix = [(label, cols.find(*cands)) for label, cands in want]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fiber Order Log"
    ws.append([label for label, _ in ix])
    for c in ws[1]:
        c.font = Font(bold=True)
    i_rep = cols.need("rep", "rep")
    i_od = cols.need("order date", "sp.order date")
    i_st = cols.find("dtr status", "spe.status")
    rows = sorted(crew, key=lambda r: (_c(cols, r, i_rep),
                                       _parse_date(_c(cols, r, i_od))
                                       or dt.date.min))
    for r in rows:
        ws.append([_c(cols, r, i) for _, i in ix])
        st = _c(cols, r, i_st).lower()
        fill = FILL.get(st)
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=fill)
    for col_cells in ws.columns:
        w = max((len(str(c.value or "")) for c in col_cells), default=8)
        ws.column_dimensions[col_cells[0].column_letter].width = min(
            max(w + 2, 9), 36)
    ws.freeze_panes = "A2"
    out = OUT_DIR / "Fiber Order Log {}.xlsx".format(
        today.strftime("%m-%d-%Y"))
    wb.save(out)
    log(f"[fiber] order log workbook: {len(rows)} line(s)")
    return out


def _activation_counts(cols, crew, today):
    i_od = cols.need("order date", "sp.order date")
    i_st = cols.find("dtr status", "spe.status")
    i_rep = cols.need("rep", "rep")
    out: Dict[str, dict] = {}

    def slot(key, bucket):
        return out.setdefault(key, {b: {"sold": 0, "act": 0}
                                    for b in BUCKETS})[bucket]

    for r in crew:
        od = _parse_date(_c(cols, r, i_od))
        if od is None:
            continue
        age = (today - od).days
        bucket = ("0-30 Day" if 0 <= age <= 30
                  else "31-60 Day" if 31 <= age <= 60 else None)
        if bucket is None:
            continue
        st = _c(cols, r, i_st).lower()
        act = st in POSTED or st == "disconnected"   # was active before it fell
        for key in (_c(cols, r, i_rep) or "(no rep)", "__TOTAL__"):
            c = slot(key, bucket)
            c["sold"] += 1
            if act:
                c["act"] += 1
    return out


def build_activation_png(cols, crew, today, log=print) -> Path:
    from automations.b2b_metrics.rep_boards import render_table_png
    from automations.vantura_churn.fill import BANDS
    counts = _activation_counts(cols, crew, today)

    def _color(bucket, rate):
        for floor, name in BANDS[_BAND_KEY[bucket]]:
            if rate >= floor:
                return name.capitalize()
        return ""

    def _cells(d):
        cells = {}
        for b in BUCKETS:
            c = d.get(b) or {}
            if not c.get("sold"):
                cells[b] = {}
                continue
            rate = c["act"] / c["sold"]
            cells[b] = {"act": str(c["sold"]), "disc": str(c["act"]),
                        "rate": f"{round(rate * 100, 1)}%",
                        "color": _color(b, rate)}
        return cells

    rows = [("Crew Total (all reps)", "", True, _cells(counts["__TOTAL__"]))]
    for rep in sorted(k for k in counts if k != "__TOTAL__"):
        rows.append((rep, "", False, _cells(counts[rep])))
    for b in BUCKETS:
        t = counts["__TOTAL__"][b]
        log(f"[fiber] CREW {b}: {t['act']}/{t['sold']}")
    return render_table_png(
        "FIBER — ACTIVATION RATES BY REP",
        f"Carlos's crew under Raf's code — {today.strftime('%B %d, %Y')} "
        "(activated/sold by sale-date age from the D2D order log; cancels "
        "count as sold; Posted = activated; total = whole crew)",
        list(BUCKETS), rows, OUT_DIR / "fiber_activation_by_rep.png")


def build_pending_png(cols, crew, today, log=print) -> Path:
    """Chase list: sold in the last 31 days, not activated, not dead."""
    from automations.box_order_log import pending_png
    from automations.box_order_log import pending as bp
    i_od = cols.need("order date", "sp.order date")
    i_st = cols.find("dtr status", "spe.status")
    i_rep = cols.need("rep", "rep")
    i_cust = cols.find("customer name")
    COLS = ("Customer", "Sale Date", "Days Waiting", "Status")

    class _Row:
        def __init__(self, cells, status, rep, sale_date, business):
            self.cells, self.status = cells, status
            self.history = ()
            self.rep, self.sale_date = rep, sale_date
            self.fields = {"Rep Name": rep, "Business Name": business}

    rows = []
    for r in crew:
        od = _parse_date(_c(cols, r, i_od))
        if od is None or (today - od).days > 31:
            continue
        st = _c(cols, r, i_st)
        sl = st.lower()
        if sl in POSTED or sl in DEAD:
            continue
        cust = _c(cols, r, i_cust)
        rows.append(_Row((cust, od.strftime("%m/%d/%Y"),
                          str((today - od).days), st or "—"),
                         sl, _c(cols, r, i_rep), od, cust))
    n = len(rows)
    work = {"today": today, "count": n,
            "title": "Fiber — pending orders, not yet activated",
            "subtitle": ("Crew lines sold in the last 31 days with no "
                         "activation yet — {} as of {}. Confirmed/Shipped/"
                         "blank statuses; Posted and dead lines excluded."
                         .format("{} line{}".format(n, bp.plural(n)) if n
                                 else "none right now",
                                 today.strftime("%B %d, %Y"))),
            "columns": COLS,
            "color_fn": lambda st, _h: {"shipped": "FFF2CC",
                                        "confirmed": "FFF2CC"}.get(st, ""),
            "sections": [{"key": "pending", "title": None, "rows": rows,
                          "reps": bp.by_rep(rows),
                          "empty_note": "Nothing pending — every crew line "
                                        "is activated or closed."}]}
    out = OUT_DIR / "fiber_pending_orders.png"
    pending_png.render(work, out)
    log(f"[fiber] pending orders: {n} line(s)")
    return out


def build_churn_rolloff_png(cols, crew, today, log=print) -> Path:
    """The 0-30 rolloff list — the log's own churn columns do the dating."""
    from automations.b2b_metrics.rep_boards import render_table_png
    i_od = cols.need("order date", "sp.order date")
    i_st = cols.find("dtr status", "spe.status")
    i_rep = cols.need("rep", "rep")
    i_cust = cols.find("customer name")
    i_fa = cols.find("dtr first available")
    i_dtd = cols.find("days to disconnect")
    i_why = cols.find("disconnect reason")
    i_bkt = cols.find("churn buckets")

    rows_out, n_disc = [], 0
    for r in crew:
        if _c(cols, r, i_st).lower() != "disconnected":
            continue
        bucket = _c(cols, r, i_bkt)
        if bucket and "0-30" not in bucket and "0 - 30" not in bucket:
            continue                     # only the 0-30 window chase list
        act = _parse_date(_c(cols, r, i_fa)) or _parse_date(_c(cols, r, i_od))
        dtd = _c(cols, r, i_dtd)
        disc_day = None
        try:
            if act and dtd:
                disc_day = act + dt.timedelta(days=int(float(dtd)))
        except ValueError:
            pass
        ages_off = (act + dt.timedelta(days=30)) if act else None
        days_left = (ages_off - today).days if ages_off else None
        n_disc += 1
        rows_out.append(((days_left if days_left is not None else 99),
                         (_c(cols, r, i_rep), "", False, {
            "Days Left": {"act": "", "disc": "",
                          "rate": str(days_left) if days_left is not None
                          else "?", "color": ""},
            "Customer": {"rate": _c(cols, r, i_cust)},
            "Activated": {"rate": act.strftime("%m/%d") if act else ""},
            "Disconnected": {"rate": disc_day.strftime("%m/%d")
                             if disc_day else ""},
            "Reason": {"rate": (_c(cols, r, i_why) or "")[:28]},
        })))
    rows_out.sort(key=lambda t: t[0])
    cols_list = ["Days Left", "Customer", "Activated", "Disconnected",
                 "Reason"]
    table = [t[1] for t in rows_out] or [
        ("No 0-30 disconnects right now", "", True,
         {c: {} for c in cols_list})]
    log(f"[fiber] churn rolloff: {n_disc} disconnected line(s) in 0-30")
    return render_table_png(
        "FIBER — 0-30 DAY ROLLOFF LIST",
        f"Carlos's crew under Raf's code — {today.strftime('%B %d, %Y')} "
        "(disconnected lines still inside the 0-30 window, soonest to age "
        "off first; dates from the log's own churn columns)",
        cols_list, table, OUT_DIR / "fiber_customer_churn.png")


def build_churn_rates_png(cols, crew, today, log=print) -> Path:
    """Crew churn per product per bucket — computed, since the D2D CHURN
    views only know offices."""
    from automations.b2b_metrics.rep_boards import render_table_png
    i_od = cols.need("order date", "sp.order date")
    i_st = cols.find("dtr status", "spe.status")
    i_prod = cols.find("product type (broken out)", "product")
    agg: Dict[str, dict] = {}
    for r in crew:
        od = _parse_date(_c(cols, r, i_od))
        if od is None:
            continue
        age = (today - od).days
        bucket = ("0-30 Day" if 0 <= age <= 30
                  else "31-60 Day" if 31 <= age <= 60 else None)
        if bucket is None:
            continue
        st = _c(cols, r, i_st).lower()
        if not (st in POSTED or st == "disconnected"):
            continue                       # churn base = activated lines
        prod = _c(cols, r, i_prod) or "(unknown)"
        for key in (prod, "__ALL__"):
            c = agg.setdefault(key, {b: {"act": 0, "disc": 0}
                                     for b in BUCKETS})[bucket]
            c["act"] += 1
            if st == "disconnected":
                c["disc"] += 1

    def _cells(d):
        cells = {}
        for b in BUCKETS:
            c = d.get(b) or {}
            if not c.get("act"):
                cells[b] = {}
                continue
            rate = c["disc"] / c["act"]
            cells[b] = {"act": str(c["act"]), "disc": str(c["disc"]),
                        "rate": f"{round(rate * 100, 1)}%",
                        "color": ("Green" if rate < 0.04 else
                                  "Yellow" if rate < 0.06 else "Red")}
        return cells

    rows = [("Crew Total (all products)", "", True,
             _cells(agg.get("__ALL__", {})))]
    for prod in sorted(k for k in agg if k != "__ALL__"):
        rows.append((prod, "", False, _cells(agg[prod])))
    return render_table_png(
        "FIBER — CHURN RATES",
        f"Carlos's crew under Raf's code — {today.strftime('%B %d, %Y')} "
        "(disconnects ÷ activated lines, by sale-date age; provisional "
        "bands green<4% yellow<6% until the Fiber comp sets real ones)",
        list(BUCKETS), rows, OUT_DIR / "fiber_churn_rates.png")


SECTIONS = [
    ("order_log", "\U0001F4C4", "Fiber Order Log", build_order_log),
    ("activation_by_rep", "\U0001F4C8", "Fiber Activation Rate by Rep",
     build_activation_png),
    ("pending_orders", "⏳", "Fiber Pending Orders", build_pending_png),
    ("customer_churn", "\U0001F43A", "Fiber Customer Churn",
     build_churn_rolloff_png),
    ("churn_rates", "\U0001F4C9", "Fiber Churn Rates",
     build_churn_rates_png),
]


# ---------------------------------------------------------------- post ----

def _state_path(today: dt.date) -> Path:
    return OUT_DIR / f"thread_{today.isoformat()}.json"


def post_thread(artifacts, today, log=print) -> None:
    from automations.shared.slack_metrics_post import _client
    client = _client()
    sp = _state_path(today)
    state = json.loads(sp.read_text()) if sp.exists() else {}
    ts = state.get("ts")
    if not ts:
        header = f"Fiber Metrics — {today.strftime('%B %d, %Y')}"
        ts = client.chat_postMessage(channel=CHANNEL[1], text=header)["ts"]
        contents = "\n".join(f"{emoji} {title}"
                             for _id, emoji, title, _fn in SECTIONS)
        client.chat_postMessage(channel=CHANNEL[1], thread_ts=ts,
                                text="In this thread:\n" + contents)
        state = {"ts": ts, "posted": []}
        log(f"[fiber] opened thread {ts} in {CHANNEL[0]}")
    for (sid, emoji, title, _fn) in SECTIONS:
        path = artifacts.get(sid)
        if not path or sid in state["posted"]:
            continue
        client.files_upload_v2(channel=CHANNEL[1], thread_ts=ts,
                               file=str(path), filename=Path(path).name,
                               initial_comment=f"{emoji} *{title}*")
        state["posted"].append(sid)
        sp.write_text(json.dumps(state))
        log(f"[fiber] posted {sid}")


def dm_previews(artifacts, user, log=print) -> None:
    from automations.shared import slack_metrics_post as smp
    for sid, emoji, title, _fn in SECTIONS:
        path = artifacts.get(sid)
        if not path:
            continue
        smp.dm_user_with_file(
            Path(path), user=user, file_name=Path(path).name,
            comment=f"{emoji} *{title}* — Fiber Metrics PREVIEW (crew only, "
                    "not posted).")
        log(f"[fiber] DM'd {sid}")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="fiber_metrics")
    ap.add_argument("--today", default=None, metavar="YYYY-MM-DD")
    ap.add_argument("--only", default=None,
                    help="comma-separated section ids")
    ap.add_argument("--dm", default=None, metavar="U...",
                    help="DM previews to this user; never touches the channel")
    ap.add_argument("--post", action="store_true",
                    help="post the live Fiber Metrics thread")
    a = ap.parse_args(argv)
    today = (dt.date.fromisoformat(a.today) if a.today else dt.date.today())
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    path = pull(today)
    cols, crew = load_crew(path)
    only = set(a.only.split(",")) if a.only else None
    artifacts = {}
    for sid, _e, _t, fn in SECTIONS:
        if only and sid not in only:
            continue
        try:
            artifacts[sid] = fn(cols, crew, today)
        except Exception as e:  # noqa: BLE001 — one board never sinks the rest
            print(f"[fiber] {sid} FAILED: {type(e).__name__}: {e}")
    print(f"[fiber] built: {sorted(artifacts)}")
    if a.dm:
        dm_previews(artifacts, a.dm)
    if a.post:
        post_thread(artifacts, today)
    if not a.dm and not a.post:
        print("[fiber] dry-run — --dm for previews, --post for the thread")
    print("=== done ===", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
